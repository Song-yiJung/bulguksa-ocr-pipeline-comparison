"""version_B step2 — Vision OCR + alignment 검증 (Gemini ↔ OCR 교차 검증).

LLM-first 패러다임(step1=Gemini)의 검증 단계.
step1의 gemini_raw와 Vision OCR 결과를 difflib SequenceMatcher로 alignment,
양방향 [?] 마킹으로 환각/누락 자동 검출.
"""
import os
import json
import time
import re
import difflib
import argparse
from pathlib import Path
from datetime import date, datetime
from dotenv import load_dotenv

PROJECT_ROOT = Path(__file__).parent.parent
load_dotenv(dotenv_path=PROJECT_ROOT / ".env")
os.environ['GOOGLE_APPLICATION_CREDENTIALS'] = str(PROJECT_ROOT / 'vision_key.json')

from google.cloud import vision
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ============================================================
# 메타데이터
# ============================================================
VERSION_TAG = "v_B_step2_vision"
VERSION_NAME = "Version B step2 — Vision OCR + alignment 검증 (양방향 [?])"

# ============================================================
# CLI 인자
# ============================================================
parser = argparse.ArgumentParser(description="step2: Vision OCR + step1 Gemini alignment 검증")
parser.add_argument('--input-dir', default=os.getenv('OCR_INPUT_DIR',
                    r'G:\jggmk2\OCR_PIPLINE_IMAGES'))
parser.add_argument('--results-dir', default=r'D:\OCR_PIPLINE\results\version_B',
                    help='step1 결과 폴더 (= step2 입출력)')
parser.add_argument('--once', action='store_true')
args = parser.parse_args()

INPUT_DIR = Path(args.input_dir)
OUTPUT_DIR = Path(args.results_dir)
LOG_XLSX = OUTPUT_DIR / 'version_B_작업내역.xlsx'
SOURCE_XLSX = INPUT_DIR / 'OCR데이터목록.xlsx'
POLL_INTERVAL = 60

# ============================================================
# Vision 클라이언트
# ============================================================
print(f"[System] Vision API 초기화...")
try:
    vision_client = vision.ImageAnnotatorClient()
    print(f"[System] 준비 완료\n")
except Exception as e:
    print(f"[Error] Vision 초기화 실패: {e}")
    exit(1)


def vision_ocr(img_path: Path) -> str:
    with open(img_path, 'rb') as f:
        content = f.read()
    image = vision.Image(content=content)
    response = vision_client.document_text_detection(image=image)
    if response.error.message:
        raise Exception(f"Vision API 오류: {response.error.message}")
    return response.full_text_annotation.text or ""


# ============================================================
# Alignment 검증 (양방향 [?] 마킹)
# ============================================================
PUNCTS = set(' \n\t|()[]{},.;:?!-_/\\\'"·　')


def cross_validate(gemini_text: str, ocr_text: str):
    """Gemini ↔ OCR alignment 후 양방향 [?] 표시 + 통계 반환.

    Edit operation:
      equal   → 두 결과 일치, 그대로 보존
      delete  → Gemini에만 있음 → [G:글자?]  (Gemini 환각 후보)
      insert  → OCR에만 있음    → [O:글자?]  (Gemini 누락 또는 OCR 환각)
      replace → 둘 다 있지만 다름 → [G:글자|O:글자?]  (둘 중 오인식)
    """
    matcher = difflib.SequenceMatcher(None, gemini_text, ocr_text, autojunk=False)
    out = []
    stats = {
        'gemini_chars': len(gemini_text),
        'ocr_chars': len(ocr_text),
        'equal_chars': 0,
        'gemini_only_chars': 0,
        'ocr_only_chars': 0,
        'replace_chars_g': 0,
        'replace_chars_o': 0,
    }
    for tag, i1, i2, j1, j2 in matcher.get_opcodes():
        g_part = gemini_text[i1:i2]
        o_part = ocr_text[j1:j2]
        if tag == 'equal':
            out.append(g_part)
            stats['equal_chars'] += len(g_part)
        elif tag == 'delete':
            stripped = ''.join(c for c in g_part if c not in PUNCTS)
            if stripped:
                out.append(f"[G:{g_part}?]")
                stats['gemini_only_chars'] += len(stripped)
            else:
                out.append(g_part)  # 구두점/공백만이면 그대로 보존
        elif tag == 'insert':
            stripped = ''.join(c for c in o_part if c not in PUNCTS)
            if stripped:
                out.append(f"[O:{o_part}?]")
                stats['ocr_only_chars'] += len(stripped)
        elif tag == 'replace':
            out.append(f"[G:{g_part}|O:{o_part}?]")
            stats['replace_chars_g'] += len(g_part)
            stats['replace_chars_o'] += len(o_part)
    return ''.join(out), stats


def count_real_chars(text: str) -> int:
    cleaned = re.sub(r'\[[GO]:[^?]*\?\]', '', text)  # [?] 마커 제거
    return len(cleaned.replace("\n", "").replace(" ", "").replace("\t", ""))


def _expand_range(s: str):
    m = re.match(r'(.+-)(\d+)~.+?(\d+)$', s)
    if m:
        p, start, end = m.group(1), int(m.group(2)), int(m.group(3))
        w = len(m.group(2))
        return {f'{p}{str(i).zfill(w)}' for i in range(start, end + 1)}
    return {s}


# ============================================================
# Excel 작업 로그
# ============================================================
def write_version_xlsx(completed_ids: set):
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    today = date.today().isoformat()

    source_rows = []
    source_headers = []
    if SOURCE_XLSX.exists():
        try:
            src_wb = openpyxl.load_workbook(SOURCE_XLSX, read_only=True, data_only=True)
            if 'Sheet1' in src_wb.sheetnames:
                ws = src_wb['Sheet1']
                rows_iter = ws.iter_rows(values_only=True)
                source_headers = list(next(rows_iter, []) or [])
                for row in rows_iter:
                    if any(c is not None and str(c).strip() for c in row):
                        source_rows.append(list(row))
            src_wb.close()
        except Exception as e:
            print(f"  [원본 엑셀 읽기 실패(무시): {str(e)[:120]}]")

    wb = openpyxl.Workbook()
    header_font = Font(name='맑은 고딕', size=11, bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='305496', end_color='305496', fill_type='solid')
    title_font = Font(name='맑은 고딕', size=14, bold=True, color='305496')
    center = Alignment(horizontal='center', vertical='center', wrap_text=True)
    left_wrap = Alignment(horizontal='left', vertical='center', wrap_text=True)
    thin = Side(border_style='thin', color='B4B4B4')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    # Sheet1: 작업현황
    ws1 = wb.active
    ws1.title = '작업현황'
    ws1['A1'] = f'{VERSION_NAME} - 작업 현황'
    ws1['A1'].font = title_font
    ws1.row_dimensions[1].height = 22
    if source_headers:
        new_headers = list(source_headers) + ['step2_B_완료일']
        for col_idx, h in enumerate(new_headers, 1):
            cell = ws1.cell(row=3, column=col_idx, value=str(h) if h is not None else '')
            cell.font = header_font; cell.fill = header_fill
            cell.alignment = center; cell.border = border
        for row_idx, row_data in enumerate(source_rows, start=4):
            for col_idx, val in enumerate(row_data, 1):
                cell = ws1.cell(row=row_idx, column=col_idx, value=val)
                cell.alignment = left_wrap; cell.border = border
            files_value = row_data[13] if len(row_data) > 13 else None
            if files_value:
                row_ids = set()
                for part in str(files_value).split('|'):
                    row_ids.update(_expand_range(part.strip()))
                if row_ids and row_ids.issubset(completed_ids):
                    cell = ws1.cell(row=row_idx, column=len(new_headers), value=today)
                    cell.alignment = center; cell.border = border
        for col_idx in range(1, len(new_headers) + 1):
            ws1.column_dimensions[get_column_letter(col_idx)].width = 16

    # Sheet2: 상세로그 (alignment 통계 포함)
    ws2 = wb.create_sheet('상세로그')
    ws2['A1'] = f'{VERSION_NAME} - 상세 로그 + alignment 통계'
    ws2['A1'].font = title_font; ws2.row_dimensions[1].height = 22
    detail_headers = ['순번', '처리일자', '권', '문건', '문서제목', '파일ID',
                      'Gemini자수', 'Vision자수', '일치자수', 'Gemini만(환각의심)', 'OCR만(누락의심)',
                      '치환자수', '상태', '경로']
    for col_idx, h in enumerate(detail_headers, 1):
        cell = ws2.cell(row=3, column=col_idx, value=h)
        cell.font = header_font; cell.fill = header_fill
        cell.alignment = center; cell.border = border

    records = []
    for f in OUTPUT_DIR.rglob('*.json'):
        if f.name.startswith('_'):
            continue
        try:
            d = json.loads(f.read_text(encoding='utf-8'))
        except Exception:
            continue
        rel = f.relative_to(OUTPUT_DIR)
        parts = rel.parts
        mod_date = datetime.fromtimestamp(f.stat().st_mtime).date().isoformat()
        if 'error' in d:
            status = '오류'
            stats = {}
        elif d.get('final_corrected', '').strip():
            status = '성공'
            stats = d.get('alignment_stats', {})
        else:
            continue
        records.append({
            'date': mod_date,
            'kwon': parts[0] if len(parts) > 0 else '',
            'mun':  parts[1] if len(parts) > 1 else '',
            'title':parts[2] if len(parts) > 2 else '',
            'file_id': f.stem,
            'g_chars': stats.get('gemini_chars', 0),
            'o_chars': stats.get('ocr_chars', 0),
            'eq': stats.get('equal_chars', 0),
            'g_only': stats.get('gemini_only_chars', 0),
            'o_only': stats.get('ocr_only_chars', 0),
            'repl': stats.get('replace_chars_g', 0) + stats.get('replace_chars_o', 0),
            'status': status,
            'path': str(rel),
        })

    records.sort(key=lambda r: (r['date'], r['path']))
    for i, r in enumerate(records, 1):
        row_idx = 3 + i
        values = [i, r['date'], r['kwon'], r['mun'], r['title'], r['file_id'],
                  r['g_chars'], r['o_chars'], r['eq'], r['g_only'], r['o_only'], r['repl'],
                  r['status'], r['path']]
        for col_idx, v in enumerate(values, 1):
            cell = ws2.cell(row=row_idx, column=col_idx, value=v)
            if col_idx in (1, 2, 6, 7, 8, 9, 10, 11, 12, 13):
                cell.alignment = center
            else:
                cell.alignment = left_wrap
            cell.border = border
            if col_idx == 13:
                if r['status'] == '오류':
                    cell.fill = PatternFill(start_color='F8CBAD', end_color='F8CBAD', fill_type='solid')
                else:
                    cell.fill = PatternFill(start_color='C6EFCE', end_color='C6EFCE', fill_type='solid')
            elif col_idx == 10 and r['g_only'] > 50:  # 환각 의심 강조
                cell.fill = PatternFill(start_color='FFE699', end_color='FFE699', fill_type='solid')

    widths = [6, 12, 22, 20, 25, 22, 10, 10, 10, 14, 12, 10, 10, 50]
    for i, w in enumerate(widths, 1):
        ws2.column_dimensions[get_column_letter(i)].width = w
    ws2.freeze_panes = 'A4'

    # Sheet3: 요약통계
    ws3 = wb.create_sheet('요약통계')
    ws3['A1'] = f'{VERSION_NAME} - 요약'
    ws3['A1'].font = title_font; ws3.row_dimensions[1].height = 22
    total = len(records)
    success = sum(1 for r in records if r['status'] == '성공')
    error_count = sum(1 for r in records if r['status'] == '오류')
    total_g_only = sum(r['g_only'] for r in records)
    total_o_only = sum(r['o_only'] for r in records)
    total_eq = sum(r['eq'] for r in records)
    summary = [
        ('파이프라인', VERSION_NAME),
        ('생성일시', datetime.now().strftime('%Y-%m-%d %H:%M:%S')),
        ('', ''),
        ('총 처리', total),
        ('  ├ 성공', success),
        ('  └ 오류', error_count),
        ('', ''),
        ('=== alignment 통계 (합산) ==='),
        ('  일치 글자', total_eq),
        ('  Gemini만 (환각 의심)', total_g_only),
        ('  OCR만 (Gemini 누락)', total_o_only),
        ('', ''),
        ('OCR 엔진 (step2)', 'Google Vision DOCUMENT_TEXT_DETECTION'),
        ('alignment', 'difflib SequenceMatcher.get_opcodes()'),
        ('마킹 형식', '[G:글자?] / [O:글자?] / [G:글자|O:글자?]'),
        ('', ''),
        ('입력 이미지', str(INPUT_DIR)),
        ('결과', str(OUTPUT_DIR)),
    ]
    for i, item in enumerate(summary, start=3):
        if isinstance(item, tuple) and len(item) == 2:
            label, value = item
        else:
            label, value = (item[0], '') if isinstance(item, tuple) else (str(item), '')
        ws3.cell(row=i, column=1, value=label).font = Font(bold=True)
        ws3.cell(row=i, column=2, value=value).alignment = left_wrap
    ws3.column_dimensions['A'].width = 28
    ws3.column_dimensions['B'].width = 60

    try:
        wb.save(LOG_XLSX)
        print(f"  [엑셀 갱신: {LOG_XLSX.name} - 총 {total}건]")
    except PermissionError:
        print(f"  [오류: {LOG_XLSX.name} 파일이 열려있어 저장 실패]")
    except Exception as e:
        print(f"  [엑셀 저장 실패(무시): {str(e)[:120]}]")


# ============================================================
# 메인
# ============================================================
def process_pipeline():
    if not OUTPUT_DIR.exists():
        print(f"[Error] step1 결과 폴더 없음: {OUTPUT_DIR}")
        return 0

    json_files = sorted(OUTPUT_DIR.rglob('*.json'))
    json_files = [f for f in json_files if not f.name.startswith('_')]

    todo = []
    for jp in json_files:
        try:
            d = json.loads(jp.read_text(encoding='utf-8'))
            if not d.get('gemini_raw', '').strip():
                continue
            if d.get('final_corrected', '').strip():
                continue  # 이미 처리됨
            todo.append(jp)
        except Exception:
            pass

    if not todo:
        return 0

    print(f"[step2] 신규 처리 대상: {len(todo)}장")
    processed = 0
    for i, jp in enumerate(todo, 1):
        try:
            d = json.loads(jp.read_text(encoding='utf-8'))
            gemini_raw = d.get('gemini_raw', '')
            rel_file_path = d.get('file_path', '')
            img_path = INPUT_DIR / rel_file_path
            if not img_path.exists():
                print(f"[Skip] 이미지 없음: {img_path}")
                continue

            print(f"[{i}/{len(todo)}] {rel_file_path}", flush=True)
            start = time.time()
            ocr_text = vision_ocr(img_path)
            final_corrected, stats = cross_validate(gemini_raw, ocr_text)
            elapsed = time.time() - start

            d['vision_raw'] = ocr_text
            d['final_corrected'] = final_corrected
            d['alignment_stats'] = stats
            d['step2_time_sec'] = round(elapsed, 1)

            jp.write_text(json.dumps(d, ensure_ascii=False, indent=2), encoding='utf-8')

            final_path = jp.with_name(f"{jp.stem}_final.txt")
            final_path.write_text(final_corrected, encoding='utf-8')

            g_only = stats['gemini_only_chars']
            o_only = stats['ocr_only_chars']
            eq = stats['equal_chars']
            print(f"  -> Gemini {len(gemini_raw)}자 / Vision {len(ocr_text)}자 / 일치 {eq}자 / "
                  f"환각의심 {g_only}자 / 누락의심 {o_only}자 ({elapsed:.1f}s)")
            processed += 1
        except Exception as e:
            print(f"  -> [Error] {type(e).__name__}: {str(e)[:200]}")
            try:
                d['step2_error'] = f"{type(e).__name__}: {str(e)[:200]}"
                jp.write_text(json.dumps(d, ensure_ascii=False, indent=2), encoding='utf-8')
            except Exception:
                pass

    if processed > 0:
        completed_ids = set()
        for f in OUTPUT_DIR.rglob('*.json'):
            if f.name.startswith('_'):
                continue
            try:
                d = json.loads(f.read_text(encoding='utf-8'))
                if d.get('final_corrected', '').strip():
                    completed_ids.add(f.stem)
            except Exception:
                pass
        write_version_xlsx(completed_ids)

    return processed


def run_daemon():
    print("=" * 60)
    print(f"[System] {VERSION_NAME}")
    print(f"[Config] 이미지 디렉토리: {INPUT_DIR}")
    print(f"[Config] step1 결과 (입출력): {OUTPUT_DIR}")
    print("=" * 60)

    if args.once:
        process_pipeline()
        return

    print(f"[System] daemon 모드 (폴링 {POLL_INTERVAL}s, Ctrl+C 종료)")
    print("-" * 60)
    try:
        while True:
            count = process_pipeline()
            if count > 0:
                print(f"[System] {count}건 처리, 대기 중...")
            time.sleep(POLL_INTERVAL)
    except KeyboardInterrupt:
        print("\n[System] 종료")


if __name__ == '__main__':
    run_daemon()
