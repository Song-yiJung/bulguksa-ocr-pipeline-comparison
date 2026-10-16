"""지시문 상세화의 '종류'별 비교 — 지시문 분류 + 종류별 모델 행동 측정.

실험 당시 분석 코드(종류별_분석.py)의 공개본이다. 계산은 바꾸지 않았고, 경로를 인자로 받게 하고
지시문 이름을 공개명으로 바꾸었다(v1=초기, v2=수정, v3=상세화). 코드 안의 줄 번호 표(V1_MAP, V3_MAP)와
변수 이름의 v3·v4b는 실험 당시 내부명으로, 각각 공개명 v2·v3을 가리킨다.

사용 예:
  python evaluation/analyze_by_type.py --sample 표본목록.csv --gold-dir 기준판독문 \
      --results 실험데이터/출력/gemini-3.8-flash --out 실험데이터/분석결과/gemini-3.8-flash

1) v1·v2·v3 지시문을 줄 단위로 8가지 종류로 분류하고 종류별 분량(글자 수)을 센다.
2) 종류마다 그 지시가 겨냥한 모델 행동을 기존 출력(실험데이터/출력)에서 측정하고,
   v1→v2, v2→v3를 페이지별 짝 비교(Wilcoxon)한다. API 호출 없음.

산출물: 실험데이터/분석결과/종류별_분석.txt, 종류별_분석.xlsx
"""
import re
from collections import Counter
from pathlib import Path

import numpy as np
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from scipy.stats import wilcoxon

import analyze as A

HERE = Path(__file__).parent
OUT = None
VERSIONS = ['v1', 'v2', 'v3']   # 공개명. 내부명 v1·v3·v4b

TYPES = ['① 맥락 제시', '② 근거 원칙', '③ 전사 절차', '④ 표기 규약',
         '⑤ 규칙 예시', '⑥ 구조 템플릿', '⑦ 내용 어휘', '⑧ 출력 형식']

# 줄 번호 → 종류. 제목 줄(예: "1. 표·목록 레이아웃 보존")과 빈 줄은 분류하지 않는다(None).
V1_MAP = {1: 0, 2: 0, 4: 1, 7: 2, 8: 6, 9: 2, 10: 2, 11: 7}
# v1 8행 '도메인 기반 오독 보정'은 "대상이 불국사임을 전제로" 맥락에서 답을 추론하게 하는 지시이므로 ⑦로 분류한다.
V3_MAP = {1: 0, 2: 0, 4: 1,
          9: 2, 10: 2, 11: 4, 12: 4, 13: 4, 14: 4, 15: 4, 16: 2,
          19: 2, 20: 4, 21: 2, 22: 2, 23: 4, 24: 4, 25: 4, 26: 2,
          29: 5, 30: 5, 31: 5, 32: 5, 33: 5, 34: 5, 35: 5, 36: 5, 37: 5, 38: 5, 39: 2,
          42: 2, 43: 1, 44: 3,
          47: 2, 48: 4, 49: 4, 50: 2,
          53: 6, 54: 6, 55: 6, 56: 6,
          59: 1, 60: 3, 61: 4, 62: 3, 63: 1,
          66: 3, 67: 4, 68: 1, 69: 2,
          72: 2, 73: 5, 74: 2,
          77: 7, 78: 2, 79: 7}
# v4b에서 v3에 없던 줄: 내용 어휘(⑦) 또는 표기 규약(④)
V4B_NEW_NOTATION = ('인장이 찍힌 자리', '미날인 슬래시', "'件名' 항목은")
HEADER = re.compile(r'^\s*(\d+\.\s|\[)')


def classify(version):
    lines = (A.PROMPT_DIR / A.PROMPT_FILES[version]).read_text(encoding='utf-8').splitlines()
    rows = []
    if version in ('v1', 'v2'):
        m = V1_MAP if version == 'v1' else V3_MAP
        for i, line in enumerate(lines, 1):
            rows.append((i, m.get(i), line))
        return rows
    v3_lines = (A.PROMPT_DIR / A.PROMPT_FILES['v2']).read_text(encoding='utf-8').splitlines()
    v3_type = {v3_lines[i - 1]: t for i, t in V3_MAP.items()}
    for i, line in enumerate(lines, 1):
        if not line.strip() or HEADER.match(line):
            t = None
        elif line in v3_type:
            t = v3_type[line]
        elif i <= 2:
            t = 0
        elif any(k in line for k in V4B_NEW_NOTATION):
            t = 3
        elif line.strip() == '- 표기 규칙:':
            t = None
        else:
            t = 6
        rows.append((i, t, line))
    return rows


def size(rows):
    c = Counter()
    for _, t, line in rows:
        if t is not None:
            c[t] += len(line.strip())
    return c


# ── 종류별 행동 지표 ─────────────────────────────────────────────
NEW_FORMS = '台亀仏国関旧芸図'   # v3 5번에 명시한 신자체 (구자체: 臺龜佛國關舊藝圖)
V3_SURVEY_FIELDS = ['名稱', '所在地', '地域、地番', '地目及地積', '工作物其ノ他', '現狀', '由來又ハ傳說',
                    '保存上必要ト認メラルル事項', '其他參考ト爲ルベキ事項']
V3_FORM_TERMS = ['未完結', '完結', '施行月日番號', '接受', '起案', '發送', '記帳', '鈐印', '校合', '淨寫', '關係', '番號', '決裁']
V3_LEXICON = ['長尾欽彌', '藤島亥治郞', '阿彌陀如來', '毘盧舍那佛', '舍利塔', '石窟庵', '多寶塔', '釋迦塔', '白雲橋',
              '青雲橋', '七寶橋', '蓮華橋', '芬皇寺', '鳳停寺', '開心寺', '淨惠寺', '栢栗寺', '馬洞里', '凡谷里']


def term_counts(terms, text):
    pat = re.compile('|'.join(sorted(map(re.escape, terms), key=len, reverse=True)))
    return Counter(pat.findall(re.sub(r'\s+', '', text)))


def restored_inserted(terms, out, gold):
    o, g = term_counts(terms, out), term_counts(terms, gold)
    ok = sum(min(n, g.get(t, 0)) for t, n in o.items())
    return ok, sum(o.values()) - ok


def metrics(out, gold):
    """한 출력의 종류별 행동 지표."""
    m = {}
    # ③ 전사 절차 — 구자체 보존: 정답에 없는 신자체 글자 수
    m['③ 신자체 오용(자)'] = sum(max(0, out.count(ch) - gold.count(ch)) for ch in NEW_FORMS)
    # ③ 전사 절차 — 반복 기호 〃: 정답의 〃를 얼마나 살렸는가(개수)
    m['③ 〃 복원(개)'] = min(out.count('〃'), gold.count('〃'))
    # ④ 표기 규약 — [?]·□·(판독불가)·(鈐印) 사용 횟수
    m['④ 불확실성 표기(회)'] = out.count('[?]') + out.count('□') + out.count('(판독불가)')
    m['④ (鈐印) 표기(회)'] = out.count('鈐印)')
    # ⑥ 구조 템플릿 — 조사서 항목·기안용지 양식 어휘: 복원/삽입
    m['⑥ 조사서 항목 복원'], m['⑥ 조사서 항목 삽입'] = restored_inserted(V3_SURVEY_FIELDS, out, gold)
    m['⑥ 양식 어휘 복원'], m['⑥ 양식 어휘 삽입'] = restored_inserted(V3_FORM_TERMS, out, gold)
    # ⑦ 내용 어휘 — v3 6번 어휘 / v4b 추가 어휘: 복원/삽입
    m['⑦ 수정 지시문 어휘 복원'], m['⑦ 수정 지시문 어휘 삽입'] = restored_inserted(V3_LEXICON, out, gold)
    m['⑦ 상세화 추가 어휘 복원'], m['⑦ 상세화 추가 어휘 삽입'] = restored_inserted(V4B_TERMS, out, gold)
    # ⑧ 출력 형식 — 정답에 한글이 없는데 출력에 한글 설명이 섞였는가, 마크다운 기호
    # 표기 규약에 따른 한글 기호((판독불가), (직인: …))는 설명 혼입이 아니므로 뺀다
    marks = re.compile(r'\(판독불가\)|\(직인[^)]*\)')
    m['⑧ 한글 혼입(자)'] = (len(re.findall(r'[가-힣]', marks.sub('', out)))
                          if not re.search(r'[가-힣]', marks.sub('', gold)) else 0)
    m['⑧ 마크다운 기호(회)'] = out.count('```') + out.count('**') + len(re.findall(r'(?m)^#+\s', out))
    return m


V4B_TERMS = []


def main():
    global V4B_TERMS, OUT
    import argparse
    ap = argparse.ArgumentParser()
    ap.add_argument('--sample', required=True)
    ap.add_argument('--gold-dir', required=True)
    ap.add_argument('--results', required=True)
    ap.add_argument('--out', required=True)
    args = ap.parse_args()
    results_dir = Path(args.results)
    A.SAMPLE, A.GOLD_DIR = args.sample, Path(args.gold_dir)
    OUT = Path(args.out)
    OUT.mkdir(parents=True, exist_ok=True)
    pattern, term_group, _ = A.build_term_matcher()
    V4B_TERMS = list(term_group)

    # 1) 지시문 분류
    cls = {v: classify(v) for v in VERSIONS}
    sizes = {v: size(cls[v]) for v in VERSIONS}

    # 2) 행동 지표 (본 분석 253장, 페이지별 3회 평균)
    pages_all = [p for p in A.load_sample() if p['main']]
    per = {v: [] for v in VERSIONS}
    pages = []
    for p in pages_all:
        outs_by_v = {v: [f.read_text(encoding='utf-8') for f in sorted((results_dir / v).glob(f"run*/{p['fid']}.txt"))]
                     for v in VERSIONS}
        if any(not o for o in outs_by_v.values()):
            continue   # 한 버전이라도 결과가 없는 면은 짝 비교에서 뺀다 (analyze.py와 같은 기준)
        pages.append(p)
        gold = p['gold_path'].read_text(encoding='utf-8')
        for v in VERSIONS:
            ms = [metrics(o, gold) for o in outs_by_v[v]]
            per[v].append({k: float(np.mean([x[k] for x in ms])) for k in ms[0]})
    print(f'분석 대상 {len(pages)}면 (결과가 빠진 면 {len(pages_all) - len(pages)}면 제외)')
    keys = list(per['v1'][0])
    # 정답지 기준값: 정답지 자체에 기호가 몇 번 있는가 (모델의 기호 사용이 옳은지 판단용)
    gold_ref = Counter()
    for p in pages:
        g = p['gold_path'].read_text(encoding='utf-8')
        gold_ref['④ 불확실성 표기(회)'] += g.count('[?]') + g.count('□') + g.count('(판독불가)')
        gold_ref['④ (鈐印) 표기(회)'] += g.count('鈐印)')
        gold_ref['③ 〃 복원(개)'] += g.count('〃')
    bases = [p['base'] for p in pages]

    def cmp(a, b, k, base=None):
        idx = [i for i, bb in enumerate(bases) if base is None or bb == base]
        x = np.array([per[a][i][k] for i in idx]); y = np.array([per[b][i][k] for i in idx]); d = y - x; nz = d[d != 0]
        p = wilcoxon(nz).pvalue if len(nz) else 1.0
        return x.sum(), y.sum(), int((d > 0).sum()), int((d < 0).sum()), p

    lines = []
    say = lines.append
    say('지시문 상세화의 종류별 비교')
    say('=' * 100)
    say('[1] 종류별 분량 (글자 수, 제목·빈 줄 제외)')
    say(f"{'종류':<12}" + ''.join(f'{v:>10}' for v in VERSIONS) + f"{'v1→v2':>10}{'v2→v3':>10}")
    for i, t in enumerate(TYPES):
        s = [sizes[v].get(i, 0) for v in VERSIONS]
        say(f'{t:<12}' + ''.join(f'{x:>10}' for x in s) + f'{s[1] - s[0]:>+10}{s[2] - s[1]:>+10}')
    tot = [sum(sizes[v].values()) for v in VERSIONS]
    say(f"{'합계':<12}" + ''.join(f'{x:>10}' for x in tot) + f'{tot[1] - tot[0]:>+10}{tot[2] - tot[1]:>+10}')
    say('')
    say('[2] 종류별 모델 행동 (본 분석 253장 합계, 페이지별 3회 평균; 증가/감소 페이지 수, Wilcoxon p)')
    for k in keys:
        say(f'{k}' + (f'   (정답지 합계 {gold_ref[k]})' if k in gold_ref else ''))
        for a, b in [('v1', 'v2'), ('v2', 'v3')]:
            xa, xb, up, dn, p = cmp(a, b, k)
            say(f'    {a}→{b:<4} {xa:8.1f} → {xb:8.1f}   증가 {up:3d} / 감소 {dn:3d}   p={p:.2g}')
    say('')
    say('[3] 삽입 지표의 정답지 바탕별 점검 (정답지가 특정 버전 출력을 닮아서 생긴 착시인지 확인)')
    for k in [k for k in keys if '삽입' in k]:
        say(f'{k}')
        for a, b in [('v1', 'v2'), ('v2', 'v3')]:
            parts = []
            for base in ('v1·v2 시기', 'v3'):   # 판독문바탕 값 (실험 당시 내부명 그대로)
                xa, xb, up, dn, p = cmp(a, b, k, base)
                parts.append(f'바탕 {base}: {xa:6.1f}→{xb:6.1f} (p={p:.2g})')
            say(f'    {a}→{b:<4} ' + '   '.join(parts))
    text = '\n'.join(lines)
    print(text)
    (OUT / '종류별_분석.txt').write_text(text, encoding='utf-8')

    # 엑셀: 줄 단위 분류(부록용) + 지표
    wb = openpyxl.Workbook()
    ws = wb.active; ws.title = '지시문 분류'
    ws.append(['버전', '줄', '종류', '내용'])
    for v in VERSIONS:
        for i, t, line in cls[v]:
            if line.strip():
                ws.append([v, i, TYPES[t] if t is not None else '(제목)', line])
    for c in ws[1]:
        c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', start_color='305496')
    for col, w in zip('ABCD', [6, 5, 14, 120]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = 'A2'
    ws2 = wb.create_sheet('종류별 분량')
    ws2.append(['종류'] + VERSIONS + ['v1→v2', 'v2→v3'])
    for i, t in enumerate(TYPES):
        s = [sizes[v].get(i, 0) for v in VERSIONS]
        ws2.append([t] + s + [s[1] - s[0], s[2] - s[1]])
    ws3 = wb.create_sheet('행동 지표')
    ws3.append(['지표', '비교', '이전 합계', '이후 합계', '증가 페이지', '감소 페이지', 'Wilcoxon p'])
    for k in keys:
        for a, b in [('v1', 'v2'), ('v2', 'v3')]:
            xa, xb, up, dn, p = cmp(a, b, k)
            ws3.append([k, f'{a}→{b}', round(xa, 1), round(xb, 1), up, dn, float(f'{p:.3g}')])
    for s in (ws2, ws3):
        for c in s[1]:
            c.font = Font(bold=True)
        s.column_dimensions['A'].width = 22
    wb.save(OUT / '종류별_분석.xlsx')
    print(f'\n저장: {OUT}')


if __name__ == '__main__':
    main()
