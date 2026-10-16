"""version_A step1 — Google Vision OCR로 1차 텍스트 추출.

A/C의 OCR-first 패러다임 중 A 쪽: 분리형 detection+recognition 파이프라인.
step2는 version_C_paddle_standard/step2_gemini.py 재사용
(--results-dir 인자로 이 OUTPUT_DIR을 가리키면 됨).
"""
import os
import json
import time
import re
from pathlib import Path
from datetime import date, datetime
from dotenv import load_dotenv

# 1. 환경 변수 + Vision 인증
PROJECT_ROOT = Path(__file__).parent.parent
load_dotenv(dotenv_path=PROJECT_ROOT / ".env")
os.environ['GOOGLE_APPLICATION_CREDENTIALS'] = str(PROJECT_ROOT / 'vision_key.json')

from google.cloud import vision
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ============================================================
# 버전 메타데이터
# ============================================================
VERSION_TAG = "v_A_step1_vision"
VERSION_NAME = "Version A step1 (Google Vision DOCUMENT_TEXT_DETECTION)"
CHANGE_REASON = (
    "version_A: OCR-first 패러다임 중 Google Vision API 단독 1차 추출. "
    "분리형 detection+recognition+grouping 파이프라인이라 세로쓰기·표 구조에선 "
    "reading order가 휴리스틱(가로쓰기 가정)으로 결정되어 깨질 수 있음. "
    "step2 Gemini가 이미지 + vision_raw 컨텍스트로 교정."
)

# ============================================================
# 경로
# ============================================================
INPUT_DIR = Path(r'G:\jggmk2\OCR_PIPLINE_IMAGES')
OUTPUT_DIR = Path(r'D:\OCR_PIPLINE\results\version_A')
SOURCE_XLSX = INPUT_DIR / 'OCR데이터목록.xlsx'
LOG_XLSX = OUTPUT_DIR / 'version_A_작업내역.xlsx'
VERSION_INFO_TXT = OUTPUT_DIR / '_version_info.txt'
POLL_INTERVAL = 60


def write_version_info():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    lines = [
        "=" * 60,
        f"Version A — {VERSION_TAG}",
        "=" * 60,
        f"생성일시: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}",
        f"실행 코드: {Path(__file__).name}",
        "",
        "[OCR 엔진]",
        "  - Google Vision API (DOCUMENT_TEXT_DETECTION)",
        "  - 분리형 detection + recognition + grouping 파이프라인",
        "",
        "[변경 사유]",
        f"  {CHANGE_REASON}",
        "",
        "[다음 단계]",
        "  step2: version_C_paddle_standard/step2_gemini.py 재사용",
        "  실행 명령:",
        f"    python version_C_paddle_standard/step2_gemini.py --results-dir {OUTPUT_DIR}",
        "",
        "[step2 환각 가드 적용]",
        "  - safety_settings = BLOCK_NONE",
        "  - <10자 스킵 (빈 페이지 차단)",
        "  - 10~50자 프롬프트 가드 (양식 추측 채우기 금지)",
        "  - SYSTEM_PROMPT 7번 sub-rule (빈 페이지 환각 방지)",
        "",
        "=" * 60,
    ]
    VERSION_INFO_TXT.write_text("\n".join(lines), encoding="utf-8")


# ============================================================
# Vision 클라이언트
# ============================================================
print(f"[System] Google Vision API 클라이언트 초기화...")
try:
    vision_client = vision.ImageAnnotatorClient()
    print(f"[System] Vision 준비 완료\n")
    write_version_info()
    print(f"[System] 버전 정보 기록: {VERSION_INFO_TXT.name}\n")
except Exception as e:
    print(f"[Error] Vision 초기화 실패: {e}")
    exit(1)


def _expand_range(s: str):
    m = re.match(r'(.+-)(\d+)~.+?(\d+)$', s)
    if m:
        p, start, end = m.group(1), int(m.group(2)), int(m.group(3))
        w = len(m.group(2))
        return {f'{p}{str(i).zfill(w)}' for i in range(start, end + 1)}
    return {s}


def extract_text_from_image(img_path: Path) -> str:
    """Google Vision DOCUMENT_TEXT_DETECTION로 텍스트 추출."""
    with open(img_path, 'rb') as f:
        content = f.read()
    image = vision.Image(content=content)
    response = vision_client.document_text_detection(image=image)
    if response.error.message:
        raise Exception(f"Vision API 오류: {response.error.message}")
    return response.full_text_annotation.text or ""


def count_real_chars(text: str) -> int:
    return len(text.replace("\n", "").replace(" ", "").replace("\t", ""))


def write_version_xlsx(completed_ids: set):
    """버전별 엑셀 작업 로그 생성/갱신."""
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    today = date.today().isoformat()

    source_rows = []
    source_headers = []
    if SOURCE_XLSX.exists():
        try:
            src_wb = openpyxl.load_workbook(SOURCE_XLSX, read_only=True, data_only=True)
            if 'Sheet1' in src_wb.sheetnames:
                ws = src_wb['Sheet1']
                rows_iter = ws.iter_rows(values_only=True)
                source_headers = list(next(rows_iter, []) or [])
                for row in rows_iter:
                    if any(c is not None and str(c).strip() for c in row):
                        source_rows.append(list(row))
            src_wb.close()
        except Exception as e:
            print(f"  [원본 엑셀 읽기 실패(무시): {str(e)[:120]}]")

    wb = openpyxl.Workbook()
    header_font = Font(name='맑은 고딕', size=11, bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='305496', end_color='305496', fill_type='solid')
    title_font = Font(name='맑은 고딕', size=14, bold=True, color='305496')
    center = Alignment(horizontal='center', vertical='center', wrap_text=True)
    left_wrap = Alignment(horizontal='left', vertical='center', wrap_text=True)
    thin = Side(border_style='thin', color='B4B4B4')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    # 시트 1: 작업현황
    ws1 = wb.active
    ws1.title = '작업현황'
    ws1['A1'] = f'{VERSION_NAME} - 작업 현황'
    ws1['A1'].font = title_font
    ws1.row_dimensions[1].height = 22

    if source_headers:
        new_headers = list(source_headers) + [f'step1_A_완료일']
        for col_idx, h in enumerate(new_headers, 1):
            cell = ws1.cell(row=3, column=col_idx, value=str(h) if h is not None else '')
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center
            cell.border = border

        for row_idx, row_data in enumerate(source_rows, start=4):
            for col_idx, val in enumerate(row_data, 1):
                cell = ws1.cell(row=row_idx, column=col_idx, value=val)
                cell.alignment = left_wrap
                cell.border = border

            files_value = row_data[13] if len(row_data) > 13 else None
            if files_value:
                row_ids = set()
                for part in str(files_value).split('|'):
                    row_ids.update(_expand_range(part.strip()))
                if row_ids and row_ids.issubset(completed_ids):
                    last_col = len(new_headers)
                    cell = ws1.cell(row=row_idx, column=last_col, value=today)
                    cell.alignment = center
                    cell.border = border

        for col_idx in range(1, len(new_headers) + 1):
            ws1.column_dimensions[get_column_letter(col_idx)].width = 16
    else:
        ws1['A3'] = '원본 OCR데이터목록.xlsx 파일을 찾을 수 없거나 Sheet1이 없습니다.'
        ws1['A3'].font = Font(italic=True, color='808080')

    # 시트 2: 상세로그
    ws2 = wb.create_sheet('상세로그')
    ws2['A1'] = f'{VERSION_NAME} - 상세 처리 로그'
    ws2['A1'].font = title_font
    ws2.row_dimensions[1].height = 22

    detail_headers = ['순번', '처리일자', '권', '문건', '문서제목',
                      '파일ID', '실제문자수', '전체길이', '처리시간(초)', '상태', '경로']
    for col_idx, h in enumerate(detail_headers, 1):
        cell = ws2.cell(row=3, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center
        cell.border = border

    records = []
    for f in OUTPUT_DIR.rglob('*.json'):
        if f.name.startswith('_') or any(pt.startswith('_') for pt in f.relative_to(OUTPUT_DIR).parts):
            continue
        try:
            d = json.loads(f.read_text(encoding='utf-8'))
        except Exception:
            continue

        rel = f.relative_to(OUTPUT_DIR)
        parts = rel.parts
        mod_date = datetime.fromtimestamp(f.stat().st_mtime).date().isoformat()

        if 'error' in d:
            status = '오류'
            real_chars = 0
            total_chars = 0
            ptime = 0
        else:
            text = d.get('vision_raw', '')
            real_chars = count_real_chars(text)
            total_chars = len(text)
            ptime = d.get('processing_time_sec', 0)
            status = '성공' if real_chars > 0 else '텍스트없음'

        records.append({
            'date': mod_date,
            'kwon': parts[0] if len(parts) > 0 else '',
            'mun':  parts[1] if len(parts) > 1 else '',
            'title':parts[2] if len(parts) > 2 else '',
            'file_id': f.stem,
            'real_chars': real_chars,
            'total_chars': total_chars,
            'ptime': ptime,
            'status': status,
            'path': str(rel),
        })

    records.sort(key=lambda r: (r['date'], r['path']))
    for i, r in enumerate(records, 1):
        row_idx = 3 + i
        values = [i, r['date'], r['kwon'], r['mun'], r['title'],
                  r['file_id'], r['real_chars'], r['total_chars'], r['ptime'], r['status'], r['path']]
        for col_idx, v in enumerate(values, 1):
            cell = ws2.cell(row=row_idx, column=col_idx, value=v)
            if col_idx in (1, 2, 6, 7, 8, 9, 10):
                cell.alignment = center
            else:
                cell.alignment = left_wrap
            cell.border = border

            if col_idx == 10:
                if r['status'] == '오류':
                    cell.fill = PatternFill(start_color='F8CBAD', end_color='F8CBAD', fill_type='solid')
                elif r['status'] == '텍스트없음':
                    cell.fill = PatternFill(start_color='FFE699', end_color='FFE699', fill_type='solid')
                else:
                    cell.fill = PatternFill(start_color='C6EFCE', end_color='C6EFCE', fill_type='solid')

    widths = [6, 12, 25, 22, 25, 22, 10, 10, 12, 12, 60]
    for i, w in enumerate(widths, 1):
        ws2.column_dimensions[get_column_letter(i)].width = w
    ws2.freeze_panes = 'A4'

    # 시트 3: 요약통계
    ws3 = wb.create_sheet('요약통계')
    ws3['A1'] = f'{VERSION_NAME} - 요약 통계'
    ws3['A1'].font = title_font
    ws3.row_dimensions[1].height = 22

    total = len(records)
    success = sum(1 for r in records if r['status'] == '성공')
    no_text = sum(1 for r in records if r['status'] == '텍스트없음')
    error_count = sum(1 for r in records if r['status'] == '오류')
    total_real_chars = sum(r['real_chars'] for r in records)
    avg_chars = (total_real_chars / success) if success > 0 else 0
    avg_time = (sum(r['ptime'] for r in records) / total) if total > 0 else 0

    summary = [
        ('파이프라인', VERSION_NAME),
        ('생성일시', datetime.now().strftime('%Y-%m-%d %H:%M:%S')),
        ('', ''),
        ('총 처리 파일 수', total),
        ('  ├ 성공', success),
        ('  ├ 텍스트 없음', no_text),
        ('  └ 오류', error_count),
        ('', ''),
        ('총 실제 문자 수', total_real_chars),
        ('성공 건당 평균 문자수', f'{avg_chars:.1f}'),
        ('평균 처리 시간 (초)', f'{avg_time:.1f}'),
        ('', ''),
        ('OCR 엔진', 'Google Vision API (DOCUMENT_TEXT_DETECTION)'),
        ('변경 사유', CHANGE_REASON),
        ('', ''),
        ('입력 디렉토리', str(INPUT_DIR)),
        ('출력 디렉토리', str(OUTPUT_DIR)),
        ('', ''),
        ('step2 실행 명령', f'python version_C_paddle_standard/step2_gemini.py --results-dir {OUTPUT_DIR}'),
    ]

    for i, (label, value) in enumerate(summary, start=3):
        ws3.cell(row=i, column=1, value=label).font = Font(bold=True)
        ws3.cell(row=i, column=2, value=value).alignment = left_wrap
    ws3.column_dimensions['A'].width = 30
    ws3.column_dimensions['B'].width = 70

    try:
        wb.save(LOG_XLSX)
        print(f"  [엑셀 작업내역 갱신: {LOG_XLSX.name} - 총 {total}건]")
    except PermissionError:
        print(f"  [오류: {LOG_XLSX.name} 파일이 열려있어 저장 실패]")
    except Exception as e:
        print(f"  [엑셀 저장 실패(무시): {str(e)[:120]}]")


def process_pipeline():
    if not INPUT_DIR.exists():
        print(f"[Error] 입력 경로를 찾을 수 없습니다. ({INPUT_DIR})")
        return 0

    valid_extensions = {'.png', '.jpg', '.jpeg', '.tif', '.tiff'}
    image_files = sorted([
        p for p in INPUT_DIR.rglob('*')
        if p.suffix.lower() in valid_extensions
        and not any(pt.startswith('_') for pt in p.relative_to(INPUT_DIR).parts)
    ])

    if not image_files:
        return 0

    todo_new = []
    todo_retry = []
    skip_done = []

    for img_path in image_files:
        rel_path = img_path.relative_to(INPUT_DIR)
        json_path = OUTPUT_DIR / rel_path.parent / f"{img_path.stem}.json"

        if not json_path.exists():
            todo_new.append(img_path)
            continue

        try:
            data = json.loads(json_path.read_text(encoding='utf-8'))
            if 'error' in data:
                todo_retry.append(img_path)
            elif 'vision_raw' in data:
                skip_done.append(img_path)
            else:
                todo_retry.append(img_path)
        except Exception:
            todo_retry.append(img_path)

    todo_total = todo_new + todo_retry
    if not todo_total:
        return 0

    print(f"[Scan] 입력 폴더 스캔 완료")
    print(f"  - 전체 이미지   : {len(image_files):>5}장")
    print(f"  - 처리 완료(스킵): {len(skip_done):>5}장")
    print(f"  - 신규 처리 대상: {len(todo_new):>5}장")
    if todo_retry:
        print(f"  - 에러 재시도   : {len(todo_retry):>5}장")
    print("-" * 60)

    processed = 0
    for i, img_path in enumerate(todo_total, 1):
        rel_path = img_path.relative_to(INPUT_DIR)
        out_folder = OUTPUT_DIR / rel_path.parent
        out_folder.mkdir(parents=True, exist_ok=True)

        base_name = img_path.stem
        json_path = out_folder / f"{base_name}.json"
        txt_path = out_folder / f"{base_name}.txt"

        is_retry = img_path in todo_retry
        marker = "[재시도]" if is_retry else "[신규]"
        print(f"{marker} [{i}/{len(todo_total)}] {rel_path}", flush=True)

        try:
            start_time = time.time()
            vision_text = extract_text_from_image(img_path)
            elapsed = time.time() - start_time

            result_data = {
                "file_path": str(rel_path),
                "version_tag": VERSION_TAG,
                "vision_raw": vision_text,
                "processing_time_sec": round(elapsed, 1),
                "gemini_corrected": "",
            }

            json_path.write_text(
                json.dumps(result_data, ensure_ascii=False, indent=2),
                encoding="utf-8"
            )

            if vision_text:
                txt_path.write_text(vision_text, encoding="utf-8")
                real_count = count_real_chars(vision_text)
                total_len = len(vision_text)
                print(f"   -> 완료: 실제 {real_count}자 / 전체 {total_len}자 ({elapsed:.1f}s)")
            else:
                print(f"   -> 경고: 인식된 텍스트 없음 ({elapsed:.1f}s)")

            processed += 1

        except Exception as e:
            print(f"   -> 오류: {str(e)[:120]}")
            error_data = {
                "file_path": str(rel_path),
                "version_tag": VERSION_TAG,
                "error": str(e),
                "last_attempt": datetime.now().isoformat()
            }
            json_path.write_text(
                json.dumps(error_data, ensure_ascii=False, indent=2),
                encoding="utf-8"
            )

    if processed > 0:
        completed_ids = set()
        for f in OUTPUT_DIR.rglob('*.json'):
            if f.name.startswith('_'):
                continue
            try:
                data = json.loads(f.read_text(encoding='utf-8'))
                if 'error' not in data and 'vision_raw' in data:
                    completed_ids.add(f.stem)
            except Exception:
                pass
        write_version_xlsx(completed_ids)

    return processed


def run_daemon():
    print(f"[System] {VERSION_NAME}")
    print(f"[System] 상시 감시 시작 (주기: {POLL_INTERVAL}초)")
    print(f"[System] 감시 대상: {INPUT_DIR}")
    print(f"[System] 결과 저장: {OUTPUT_DIR}")
    print(f"[System] step2 명령:")
    print(f"  python version_C_paddle_standard/step2_gemini.py --results-dir {OUTPUT_DIR}")
    print(f"[System] 종료: Ctrl + C\n" + "-" * 60)

    try:
        while True:
            count = process_pipeline()
            if count > 0:
                print("-" * 60)
                print(f"[System] 대기 중... 신규 폴더/파일이 추가되면 자동 처리")
            time.sleep(POLL_INTERVAL)
    except KeyboardInterrupt:
        print("\n[System] 사용자 요청으로 종료합니다.")


if __name__ == '__main__':
    run_daemon()
